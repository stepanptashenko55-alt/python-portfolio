"""Turn a raw sales export (CSV) into a ready-to-send Excel report with summaries and charts.

Usage:
    python report.py sample_data/sales.csv --output sales_report.xlsx
"""

import argparse
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

REQUIRED_COLUMNS = {"date", "product", "region", "quantity", "unit_price"}
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")


def load_sales(path: str) -> pd.DataFrame:
    df = pd.read_csv(path, parse_dates=["date"])
    missing = REQUIRED_COLUMNS - set(df.columns)
    if missing:
        raise ValueError(f"CSV is missing columns: {', '.join(sorted(missing))}")
    df = df.dropna(subset=list(REQUIRED_COLUMNS))
    df["revenue"] = df["quantity"] * df["unit_price"]
    df["month"] = df["date"].dt.to_period("M").astype(str)
    return df


def build_summaries(df: pd.DataFrame) -> dict[str, pd.DataFrame]:
    by_month = df.groupby("month", as_index=False).agg(orders=("revenue", "size"), revenue=("revenue", "sum"))
    by_product = (df.groupby("product", as_index=False)
                    .agg(units=("quantity", "sum"), revenue=("revenue", "sum"))
                    .sort_values("revenue", ascending=False))
    by_region = (df.pivot_table(index="region", columns="month", values="revenue", aggfunc="sum", fill_value=0)
                   .reset_index())
    kpis = pd.DataFrame({
        "metric": ["Total revenue", "Orders", "Units sold", "Average order value", "Best product", "Best region"],
        "value": [
            round(df["revenue"].sum(), 2),
            len(df),
            int(df["quantity"].sum()),
            round(df["revenue"].mean(), 2),
            by_product.iloc[0]["product"],
            df.groupby("region")["revenue"].sum().idxmax(),
        ],
    })
    return {"Summary": kpis, "By month": by_month, "By product": by_product, "By region": by_region}


def style_sheet(ws) -> None:
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center")
    for col in ws.columns:
        width = max(len(str(c.value)) if c.value is not None else 0 for c in col)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max(width + 3, 12), 45)
        for c in col[1:]:
            if isinstance(c.value, float):
                c.number_format = "#,##0.00"
    ws.freeze_panes = "A2"


def add_charts(wb) -> None:
    ws = wb["By month"]
    n = ws.max_row
    line = LineChart()
    line.title, line.y_axis.title, line.height, line.width = "Revenue by month", "Revenue", 8, 18
    line.add_data(Reference(ws, min_col=3, min_row=1, max_row=n), titles_from_data=True)
    line.set_categories(Reference(ws, min_col=1, min_row=2, max_row=n))
    ws.add_chart(line, "E2")

    ws = wb["By product"]
    n = ws.max_row
    bar = BarChart()
    bar.type, bar.title, bar.height, bar.width = "bar", "Revenue by product", 8, 18
    bar.add_data(Reference(ws, min_col=3, min_row=1, max_row=n), titles_from_data=True)
    bar.set_categories(Reference(ws, min_col=1, min_row=2, max_row=n))
    ws.add_chart(bar, "E2")


def build_report(csv_path: str, output: str) -> None:
    df = load_sales(csv_path)
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        for name, table in build_summaries(df).items():
            table.to_excel(writer, sheet_name=name, index=False)
        df.drop(columns="month").to_excel(writer, sheet_name="Raw data", index=False)

    wb = load_workbook(output)
    for ws in wb.worksheets:
        style_sheet(ws)
    add_charts(wb)
    wb.save(output)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("csv", help="sales CSV with columns: date, product, region, quantity, unit_price")
    parser.add_argument("--output", default="sales_report.xlsx")
    args = parser.parse_args()
    build_report(args.csv, args.output)
    print(f"Report saved to {Path(args.output).resolve()}")


if __name__ == "__main__":
    main()
