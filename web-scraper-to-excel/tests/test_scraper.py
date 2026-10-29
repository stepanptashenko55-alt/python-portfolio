from pathlib import Path

from openpyxl import load_workbook

import scraper

FIXTURE = Path(__file__).parent / "fixtures" / "page-1.html"
PAGE_URL = "https://books.toscrape.com/catalogue/page-1.html"


def test_parse_page():
    items = scraper.parse_page(FIXTURE.read_text(encoding="utf-8"), PAGE_URL)
    assert len(items) == 2
    assert items[0] == {
        "title": "A Light in the Attic",
        "price": 51.77,
        "rating": 3,
        "in_stock": True,
        "url": "https://books.toscrape.com/catalogue/a-light-in-the-attic_1000/index.html",
    }
    assert items[1]["rating"] == 1
    assert items[1]["in_stock"] is False


def test_save_excel_and_csv(tmp_path):
    items = scraper.parse_page(FIXTURE.read_text(encoding="utf-8"), PAGE_URL)
    xlsx = tmp_path / "out.xlsx"
    scraper.save_excel(items, str(xlsx))
    scraper.save_csv(items, str(tmp_path / "out.csv"))

    ws = load_workbook(xlsx).active
    assert ws["A1"].value == "Title"
    assert ws["B2"].value == 51.77
    assert ws.max_row == 3
    assert (tmp_path / "out.csv").read_text(encoding="utf-8-sig").startswith("title,price")
