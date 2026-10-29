"""Scrape a product catalog (books.toscrape.com demo shop) into a formatted Excel file and CSV.

Usage:
    python scraper.py --pages 5 --output books.xlsx
"""

import argparse
import csv
import logging
import time
from urllib.parse import urljoin

import requests
from bs4 import BeautifulSoup
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

BASE_URL = "https://books.toscrape.com/catalogue/page-{}.html"
RATINGS = {"One": 1, "Two": 2, "Three": 3, "Four": 4, "Five": 5}
HEADERS = {"User-Agent": "Mozilla/5.0 (compatible; PortfolioScraper/1.0)"}

log = logging.getLogger("scraper")


def fetch(session: requests.Session, url: str, retries: int = 3) -> str | None:
    """GET a page with retries and backoff. Returns None on 404 (end of pagination)."""
    for attempt in range(1, retries + 1):
        try:
            resp = session.get(url, headers=HEADERS, timeout=15)
            if resp.status_code == 404:
                return None
            resp.raise_for_status()
            resp.encoding = "utf-8"
            return resp.text
        except requests.RequestException as exc:
            log.warning("Attempt %d/%d failed for %s: %s", attempt, retries, url, exc)
            time.sleep(2 ** attempt)
    raise RuntimeError(f"Could not download {url}")


def parse_page(html: str, page_url: str) -> list[dict]:
    soup = BeautifulSoup(html, "html.parser")
    items = []
    for card in soup.select("article.product_pod"):
        link = card.select_one("h3 a")
        rating_class = card.select_one("p.star-rating")["class"]
        price_text = card.select_one("p.price_color").get_text(strip=True)
        items.append({
            "title": link["title"],
            "price": float("".join(ch for ch in price_text if ch.isdigit() or ch == ".")),
            "rating": next((RATINGS[c] for c in rating_class if c in RATINGS), None),
            "in_stock": "In stock" in card.select_one("p.availability").get_text(),
            "url": urljoin(page_url, link["href"]),
        })
    return items


def scrape(max_pages: int, delay: float = 1.0) -> list[dict]:
    results = []
    with requests.Session() as session:
        for page in range(1, max_pages + 1):
            url = BASE_URL.format(page)
            html = fetch(session, url)
            if html is None:
                log.info("No more pages after %d", page - 1)
                break
            items = parse_page(html, url)
            results.extend(items)
            log.info("Page %d: %d items (total %d)", page, len(items), len(results))
            time.sleep(delay)  # be polite to the server
    return results


def save_excel(rows: list[dict], path: str) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Products"
    columns = ["Title", "Price (£)", "Rating", "In stock", "URL"]
    ws.append(columns)
    for r in rows:
        ws.append([r["title"], r["price"], r["rating"], "Yes" if r["in_stock"] else "No", r["url"]])

    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
    for row in ws.iter_rows(min_row=2, min_col=2, max_col=2):
        row[0].number_format = "0.00"
    for i, width in enumerate([60, 12, 10, 10, 70], start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(path)


def save_csv(rows: list[dict], path: str) -> None:
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        writer = csv.DictWriter(f, fieldnames=rows[0].keys())
        writer.writeheader()
        writer.writerows(rows)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--pages", type=int, default=50, help="max pages to scrape (default: all 50)")
    parser.add_argument("--output", default="products.xlsx", help="Excel output file")
    parser.add_argument("--delay", type=float, default=1.0, help="seconds between requests")
    args = parser.parse_args()

    logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(message)s")
    rows = scrape(args.pages, args.delay)
    if not rows:
        log.error("Nothing scraped")
        return
    save_excel(rows, args.output)
    save_csv(rows, args.output.rsplit(".", 1)[0] + ".csv")
    log.info("Saved %d products to %s", len(rows), args.output)


if __name__ == "__main__":
    main()
